from datetime import datetime, timedelta, timezone
from pathlib import Path
import json, os, re, secrets, sys

from fastapi import FastAPI, Depends, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from jose import jwt, JWTError
from pydantic_settings import BaseSettings, SettingsConfigDict
from pwdlib import PasswordHash
from pydantic import BaseModel
from sqlalchemy import create_engine, String, Text, ForeignKey, DateTime, select
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker, Session
from .llm import LLMError, OpenAICompatibleProvider

class Settings(BaseSettings):
    database_url: str = "sqlite:///./assistant.db"
    jwt_secret: str = "development-only-change-me"
    llm_base_url: str = ""
    llm_api_key: str = ""
    llm_model: str = ""
    cors_origins: str = "http://localhost:3000"
    model_config = SettingsConfigDict(
        env_file=Path(__file__).resolve().parents[1] / ".env",
        env_file_encoding="utf-8",
        extra="ignore",
    )
settings = Settings()
engine = create_engine(settings.database_url, connect_args={"check_same_thread": False} if settings.database_url.startswith("sqlite") else {})
SessionLocal = sessionmaker(engine, expire_on_commit=False)
class Base(DeclarativeBase): pass
class User(Base):
    __tablename__="users"
    id: Mapped[int]=mapped_column(primary_key=True)
    email: Mapped[str]=mapped_column(String(255), unique=True, index=True)
    password_hash: Mapped[str]=mapped_column(String(255))
class Conversation(Base):
    __tablename__="conversations"
    id: Mapped[int]=mapped_column(primary_key=True)
    user_id: Mapped[int]=mapped_column(ForeignKey("users.id"), index=True)
    title: Mapped[str]=mapped_column(String(200), default="New conversation")
    created_at: Mapped[datetime]=mapped_column(DateTime, default=datetime.utcnow)
class Message(Base):
    __tablename__="messages"
    id: Mapped[int]=mapped_column(primary_key=True)
    conversation_id: Mapped[int]=mapped_column(ForeignKey("conversations.id"), index=True)
    role: Mapped[str]=mapped_column(String(20))
    content: Mapped[str]=mapped_column(Text)
    created_at: Mapped[datetime]=mapped_column(DateTime, default=datetime.utcnow)
class Document(Base):
    __tablename__="documents"
    id: Mapped[int]=mapped_column(primary_key=True)
    user_id: Mapped[int]=mapped_column(ForeignKey("users.id"), index=True)
    filename: Mapped[str]=mapped_column(String(255))
    content: Mapped[str]=mapped_column(Text)
Base.metadata.create_all(engine)
passwords=PasswordHash.recommended(); oauth=OAuth2PasswordBearer(tokenUrl="/api/auth/login")
def db():
    s=SessionLocal()
    try: yield s
    finally: s.close()
def current_user(token: str=Depends(oauth), s: Session=Depends(db)):
    try: uid=int(jwt.decode(token, settings.jwt_secret, algorithms=["HS256"])["sub"])
    except (JWTError, KeyError, ValueError): raise HTTPException(401,"Invalid authentication token")
    user=s.get(User,uid)
    if not user: raise HTTPException(401,"User not found")
    return user
def token_for(user): return jwt.encode({"sub":str(user.id),"exp":datetime.now(timezone.utc)+timedelta(hours=12)},settings.jwt_secret,algorithm="HS256")
app=FastAPI(title="MCP-Powered AI Assistant", version="0.1.0")
app.add_middleware(CORSMiddleware,allow_origins=[x.strip() for x in settings.cors_origins.split(",")],allow_credentials=True,allow_methods=["*"],allow_headers=["*"])
@app.get('/api/health')
def health(): return {"status":"ok","llm_configured":bool(settings.llm_base_url and settings.llm_api_key and settings.llm_model)}
@app.post('/api/auth/register',status_code=201)
def register(data: OAuth2PasswordRequestForm=Depends(),s:Session=Depends(db)):
    if s.scalar(select(User).where(User.email==data.username)): raise HTTPException(409,"Email already registered")
    u=User(email=data.username.lower(),password_hash=passwords.hash(data.password)); s.add(u); s.commit(); s.refresh(u)
    return {"access_token":token_for(u),"token_type":"bearer","user":{"id":u.id,"email":u.email}}
@app.post('/api/auth/login')
def login(data: OAuth2PasswordRequestForm=Depends(),s:Session=Depends(db)):
    u=s.scalar(select(User).where(User.email==data.username.lower()))
    if not u or not passwords.verify(data.password,u.password_hash): raise HTTPException(401,"Incorrect email or password")
    return {"access_token":token_for(u),"token_type":"bearer","user":{"id":u.id,"email":u.email}}
@app.get('/api/auth/me')
def me(u:User=Depends(current_user)): return {"id":u.id,"email":u.email}
@app.post('/api/conversations',status_code=201)
def new_conversation(u:User=Depends(current_user),s:Session=Depends(db)):
    c=Conversation(user_id=u.id); s.add(c); s.commit(); s.refresh(c); return {"id":c.id,"title":c.title}
@app.get('/api/conversations')
def conversations(u:User=Depends(current_user),s:Session=Depends(db)):
    return [{"id":c.id,"title":c.title,"created_at":c.created_at.isoformat()} for c in s.scalars(select(Conversation).where(Conversation.user_id==u.id).order_by(Conversation.id.desc()))]
@app.get('/api/conversations/{cid}')
def get_conversation(cid:int,u:User=Depends(current_user),s:Session=Depends(db)):
    c=s.get(Conversation,cid)
    if not c or c.user_id!=u.id: raise HTTPException(404,"Conversation not found")
    return {"id":c.id,"title":c.title,"messages":[{"role":m.role,"content":m.content} for m in s.scalars(select(Message).where(Message.conversation_id==cid).order_by(Message.id))]}
@app.delete('/api/conversations/{cid}',status_code=204)
def delete_conversation(cid:int,u:User=Depends(current_user),s:Session=Depends(db)):
    c=s.get(Conversation,cid)
    if not c or c.user_id!=u.id: raise HTTPException(404,"Conversation not found")
    for m in s.scalars(select(Message).where(Message.conversation_id==cid)): s.delete(m)
    s.delete(c); s.commit()
class ChatIn(BaseModel): conversation_id:int; message:str
async def call_llm(history):
    provider = OpenAICompatibleProvider(
        settings.llm_base_url,
        settings.llm_api_key,
        settings.llm_model,
    )
    return await provider.complete(history)
@app.post('/api/chat')
async def chat(data:ChatIn,u:User=Depends(current_user),s:Session=Depends(db)):
    c=s.get(Conversation,data.conversation_id)
    if not c or c.user_id!=u.id: raise HTTPException(404,"Conversation not found")
    text=data.message.strip()
    if not text or len(text)>12000: raise HTTPException(422,"Message must contain 1–12000 characters")
    prior=list(s.scalars(select(Message).where(Message.conversation_id==c.id).order_by(Message.id)))
    s.add(Message(conversation_id=c.id,role="user",content=text)); s.commit()
    history=[{"role":"system","content":"You are a helpful assistant. Be clear and concise."}]+[{"role":m.role,"content":m.content} for m in prior]+[{"role":"user","content":text}]
    try: answer=await call_llm(history)
    except LLMError as e: raise HTTPException(e.status_code,e.detail) from e
    s.add(Message(conversation_id=c.id,role="assistant",content=answer)); s.commit()
    return {"answer":answer,"tools_used":[]}
@app.post('/api/documents/upload',status_code=201)
async def upload(file:UploadFile=File(...),u:User=Depends(current_user),s:Session=Depends(db)):
    ext=Path(file.filename or '').suffix.lower()
    if ext not in {'.txt','.md','.pdf','.docx'}: raise HTTPException(415,"Supported formats: PDF, TXT, Markdown, DOCX")
    raw=await file.read(8*1024*1024+1)
    if len(raw)>8*1024*1024: raise HTTPException(413,"Maximum file size is 8 MiB")
    try:
        if ext in {'.txt','.md'}: content=raw.decode('utf-8-sig')
        elif ext=='.pdf':
            from pypdf import PdfReader
            import io
            content='\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(raw)).pages)
        else:
            from docx import Document as Docx
            import io
            content='\n'.join(p.text for p in Docx(io.BytesIO(raw)).paragraphs)
    except Exception as e: raise HTTPException(422,"Could not extract text from file") from e
    d=Document(user_id=u.id,filename=Path(file.filename or 'document').name,content=content[:2_000_000]); s.add(d); s.commit(); s.refresh(d)
    return {"id":d.id,"filename":d.filename,"characters":len(d.content)}
@app.get('/api/documents')
def documents(u:User=Depends(current_user),s:Session=Depends(db)): return [{"id":d.id,"filename":d.filename} for d in s.scalars(select(Document).where(Document.user_id==u.id))]
@app.post('/api/documents/{did}/search')
def search_document(did:int,q:str,u:User=Depends(current_user),s:Session=Depends(db)):
    d=s.get(Document,did)
    if not d or d.user_id!=u.id: raise HTTPException(404,"Document not found")
    terms=[x for x in re.findall(r'\w+',q.lower()) if x]
    chunks=[x.strip() for x in re.split(r'(?<=[.!?])\s+|\n+',d.content) if x.strip()]
    found=[{"text":c,"source":d.filename} for c in chunks if any(t in c.lower() for t in terms)][:10]
    return {"results":found}
@app.delete('/api/documents/{did}',status_code=204)
def delete_document(did:int,u:User=Depends(current_user),s:Session=Depends(db)):
    d=s.get(Document,did)
    if not d or d.user_id!=u.id: raise HTTPException(404,"Document not found")
    s.delete(d); s.commit()
@app.get('/api/mcp/tools')
async def mcp_tools():
    """Discover tools using the MCP Python SDK against the bundled stdio server."""
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    server=Path(__file__).resolve().parents[2]/'mcp_servers'/'demo_server.py'
    params=StdioServerParameters(command=os.getenv('MCP_SERVER_COMMAND',sys.executable),args=[str(server)])
    try:
        async with stdio_client(params) as (read,write):
            async with ClientSession(read,write) as session:
                await session.initialize(); result=await session.list_tools()
                return {"server":"demo-stdio","connected":True,"tools":[{"name":t.name,"description":t.description,"input_schema":t.inputSchema} for t in result.tools]}
    except Exception as e: raise HTTPException(503,"MCP server unavailable; install backend requirements") from e
class ExecuteIn(BaseModel): name:str; arguments:dict
@app.post('/api/mcp/execute')
async def execute_tool(data:ExecuteIn,u:User=Depends(current_user)):
    if data.name not in {'get_current_time','calculate'}: raise HTTPException(403,"Tool is not permitted")
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    server=Path(__file__).resolve().parents[2]/'mcp_servers'/'demo_server.py'
    try:
        async with stdio_client(StdioServerParameters(command=os.getenv('MCP_SERVER_COMMAND',sys.executable),args=[str(server)])) as (read,write):
            async with ClientSession(read,write) as session:
                await session.initialize(); result=await session.call_tool(data.name,data.arguments)
    except Exception as e: raise HTTPException(502,"MCP tool execution failed") from e
    if result.isError: raise HTTPException(422,"MCP tool rejected the supplied arguments")
    return {"name":data.name,"result":[c.model_dump() for c in result.content],"is_error":False}
